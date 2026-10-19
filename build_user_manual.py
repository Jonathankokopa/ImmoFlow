from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_LINE_SPACING

OUT = 'docs/Real_Estate_Management_System_User_Manual.docx'
BLUE = '163A5F'
MID = '356A9A'
PALE = 'EAF2F8'
GRAY = 'F3F5F7'
GRID = 'D9D9D9'

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_border(cell, color=GRID):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag); borders.append(element)
        element.set(qn('w:val'), 'single'); element.set(qn('w:sz'), '6'); element.set(qn('w:color'), color)

def set_cell_margin(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    mar = tcPr.first_child_found_in('w:tcMar')
    if mar is None:
        mar = OxmlElement('w:tcMar'); tcPr.append(mar)
    for side, value in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = mar.find(qn('w:' + side))
        if node is None:
            node = OxmlElement('w:' + side); mar.append(node)
        node.set(qn('w:w'), str(value)); node.set(qn('w:type'), 'dxa')

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr(); rep = OxmlElement('w:tblHeader'); rep.set(qn('w:val'), 'true'); trPr.append(rep)

def set_col_width(cell, width):
    cell.width = Inches(width)
    tcPr = cell._tc.get_or_add_tcPr(); tcW = tcPr.find(qn('w:tcW'))
    tcW.set(qn('w:w'), str(int(width * 1440))); tcW.set(qn('w:type'), 'dxa')

def set_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = 'Aptos'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic
    if color: run.font.color.rgb = RGBColor.from_string(color)

def add_text(p, text, size=10.5, bold=False, color='000000', italic=False):
    r = p.add_run(text); set_font(r, size, bold, color, italic); return r

def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.12
    add_text(p, text, 10.25)
    return p

def add_number(doc, text):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.space_after = Pt(3); p.paragraph_format.line_spacing = 1.12
    add_text(p, text, 10.25)
    return p

def add_heading(doc, text, level=1):
    p = doc.add_heading('', level=level)
    p.paragraph_format.space_before = Pt(13 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(5)
    add_text(p, text, 15 if level == 1 else 12, True, '000000')
    return p

def add_para(doc, text, bold_lead=None):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6); p.paragraph_format.line_spacing = 1.15
    if bold_lead and text.startswith(bold_lead):
        add_text(p, bold_lead, 10.5, True)
        add_text(p, text[len(bold_lead):], 10.5)
    else: add_text(p, text, 10.5)
    return p

def add_table(doc, headers, rows, widths):
    tbl = doc.add_table(rows=1, cols=len(headers)); tbl.alignment = WD_TABLE_ALIGNMENT.CENTER; tbl.style = 'Table Grid'
    tbl.autofit = False
    hdr = tbl.rows[0]; set_repeat_table_header(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]; set_col_width(cell, widths[i]); set_cell_shading(cell, BLUE); set_cell_border(cell); set_cell_margin(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_text(p, text, 9.5, True, 'FFFFFF')
    for n, row in enumerate(rows):
        cells = tbl.add_row().cells
        for i, text in enumerate(row):
            cell = cells[i]; set_col_width(cell, widths[i]); set_cell_border(cell); set_cell_margin(cell)
            if n % 2 == 1: set_cell_shading(cell, PALE)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.05
            add_text(p, str(text), 9.25)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl

doc = Document()
sec = doc.sections[0]; sec.top_margin = Inches(.72); sec.bottom_margin = Inches(.7); sec.left_margin = Inches(.78); sec.right_margin = Inches(.78)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos'); styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos'); styles['Normal'].font.size = Pt(10.5)
for s in ['Title','Heading 1','Heading 2','List Bullet','List Bullet 2','List Number']:
    styles[s].font.name = 'Aptos'; styles[s]._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos'); styles[s]._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')

# Cover
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(120); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_text(p, 'REAL ESTATE MANAGEMENT SYSTEM', 25, True, BLUE)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(10)
add_text(p, 'User Manual', 20, False, '000000')
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(32)
add_text(p, 'A practical guide for administrators, property managers, accountants, owners, staff, and tenants', 11, False, '3D4B58')
t = doc.add_table(rows=3, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
for i, (a,b) in enumerate([('Document version','1.0'),('Prepared for','System users'),('Published','October 2026')]):
    for j, text in enumerate([a,b]):
        c=t.cell(i,j); set_col_width(c, 2.2 if j==0 else 2.6); set_cell_border(c); set_cell_margin(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if j==0: set_cell_shading(c, PALE)
        p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.LEFT; add_text(p,text,10, j==0, BLUE if j==0 else '000000')
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

add_heading(doc, 'About this manual')
add_para(doc, 'This manual explains how to use the Real Estate Management System to manage properties, units, tenants, leases, rent invoices, payments, maintenance, reports, and self-service portal activities. Use the sections relevant to your role; menus and actions shown to you depend on your access rights.')
add_heading(doc, 'Who should use it', 2)
add_table(doc, ['User group','Typical responsibilities'], [
    ('Administrator','Sets up users, portfolio records, finance, maintenance, reports, and access.'),
    ('Property manager','Manages properties, units, tenants, leases, maintenance, invoices, and operational reporting.'),
    ('Accountant','Records and reviews invoices and payments; produces financial reports.'),
    ('Owner','Views and manages the portfolio and finance records assigned to them.'),
    ('Reception staff','Maintains day-to-day portfolio, tenant, and lease records where permitted.'),
    ('Maintenance staff','Works with maintenance categories, requests, and expenses.'),
    ('Tenant','Uses the tenant portal for tenancy details, payment history, maintenance requests, and profile settings.')], [1.55,5.0])
add_heading(doc, 'How the system is organised', 2)
add_para(doc, 'Most staff workflows follow this sequence: create the portfolio records, add people, establish an active lease, bill the lease, record or collect payments, then monitor maintenance and reporting. The system protects financial history by restricting changes to records that already have invoices or payments.')
add_heading(doc, 'Before you start', 2)
for item in ['Use a supported modern browser and keep it up to date.', 'Have the property, unit, tenant, lease, and payment details ready before entering records.', 'Use your organisation’s approved process for personal documents and payment information.', 'If a menu is missing, your role may not have permission. Contact an administrator rather than sharing logins.']:
    add_bullet(doc,item)

add_heading(doc, 'Quick start workflow')
add_table(doc, ['Step','Action','Where to work'], [
    ('1','Create property types, owners, properties, and units.','Property Types, Owners, Properties, Units'),
    ('2','Add tenants and create their portal accounts when needed.','Tenants'),
    ('3','Create an active lease linking a tenant to a unit.','Leases'),
    ('4','Create the rent invoice for the active lease and billing period.','Rent Invoices'),
    ('5','Record a manual payment or initiate an M-Pesa payment for an eligible KES invoice.','Payments or Rent Invoice'),
    ('6','Review receipts, maintenance activity, and management reports.','Payments, Maintenance, Reports')], [0.55,3.45,2.55])
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

add_heading(doc, '1  Signing in and everyday navigation')
add_heading(doc, 'Sign in securely', 2)
for item in ['Open the system sign-in page and enter your registered email address and password.', 'Complete email verification if prompted. Verified access is required for the protected application areas.', 'Use the password reset link if you cannot sign in. Do not reuse or share passwords.', 'Select Logout from the user menu when finishing, especially on a shared device.']:
    add_number(doc,item)
add_heading(doc, 'Use the dashboard and menu', 2)
add_para(doc, 'After signing in, the dashboard gives an at-a-glance view of your accessible work. Use the navigation menu to open a module. List pages normally let you open an existing item, create a new one, or use the available filters and action buttons. The notification bell shows recent alerts; open an alert to mark it read and go to the related work.')
add_heading(doc, 'Personal settings', 2)
add_para(doc, 'Open Settings to choose the application language (English, Swahili, or French), appearance (light, dark, or system), and enter your preferred timezone and currency. The application saves language and appearance preferences for your account. Use Profile to update your account details and change your password.')
add_heading(doc, 'Get help', 2)
add_para(doc, 'Open Support to submit a request. Choose the category, provide a specific subject, and include enough detail for someone to reproduce the issue. If AI Help is available in your environment, it can provide navigation guidance; it does not change records or replace a support request for account, billing, or technical problems.')

add_heading(doc, '2  Set up the portfolio')
add_heading(doc, 'Property types and owners', 2)
add_para(doc, 'Administrators and property managers can create property types before adding properties. Create owner records for the people or organisations connected to the portfolio. Administrators can also create an owner portal account where that feature is required.')
add_heading(doc, 'Create a property', 2)
for item in ['Open Properties and select the action to create a property.', 'Enter the property details and select the applicable type and owner where shown.', 'Save the property, then open it to review the details.', 'Add property images when needed. Set the preferred image as the cover image and remove images that should no longer be stored.']:
    add_number(doc,item)
add_heading(doc, 'Add units', 2)
add_para(doc, 'Open Units and create each rentable or managed unit. Link it to the correct property and complete the unit information requested by the form. Accurate unit records are essential because leases, maintenance, and tenant occupancy depend on this link.')
add_para(doc, 'Good practice: complete the portfolio before entering leases. Confirm spelling, ownership, and unit identifiers now; correcting linked records later can be restricted to preserve history.', 'Good practice:')

add_heading(doc, '3  Manage tenants and leases')
add_heading(doc, 'Create and maintain tenant records', 2)
for item in ['Open Tenants and select Create.', 'Enter the tenant’s contact and identification details requested by the form.', 'Save the tenant and confirm the information on the tenant profile.', 'From the tenant profile, create a portal account or toggle portal access when the tenant should use self-service. Portal accounts are intentionally created from the tenant record, not from User Management.']:
    add_number(doc,item)
add_heading(doc, 'Create a lease', 2)
add_para(doc, 'Open Leases and create a lease for an eligible unit and tenant. Enter the dates, monthly rent, currency, and any other required agreement details. Only active leases with a valid monthly rent can receive rent invoices.')
add_heading(doc, 'Lease changes and checkout', 2)
add_para(doc, 'Use the lease checkout action when a tenancy ends, following your organisation’s process. Be careful when editing a lease after invoices exist: the system may prevent changing the unit, tenant, or other history-sensitive details. A lease with rent invoices cannot be deleted; retain the financial trail and use the appropriate operational process instead.')
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

add_heading(doc, '4  Invoice rent and manage payments')
add_heading(doc, 'Create a rent invoice', 2)
for item in ['Open Rent Invoices and select Create.', 'Choose the active lease. Confirm the tenant, unit, currency, and monthly rent before continuing.', 'Enter the invoice date, due date, billing period, amount, and notes as required.', 'Save. The system generates the invoice number and checks that another invoice does not already exist for the same lease and billing period.', 'Review the saved invoice. Any available tenant credit in the matching currency may be applied automatically.']:
    add_number(doc,item)
add_para(doc, 'Invoice control: do not create a second invoice to correct an amount without first checking the original. Once payment or credit activity exists, important invoice fields are protected. A cancelled invoice cannot be edited, and an invoice with payment activity cannot be cancelled or deleted until the activity is reversed or refunded through your approved process.', 'Invoice control:')
add_heading(doc, 'Record a manual or internal payment', 2)
for item in ['Open Payments and choose Create.', 'Select the relevant unpaid invoice, then enter the payment details requested by the form.', 'Verify the amount, date, method, and reference before saving.', 'Open the payment record to review the allocation and status.', 'Use the receipt action to view or download the generated receipt. Receipts can also be verified with their verification link or token.']:
    add_number(doc,item)
add_heading(doc, 'Initiate an M-Pesa payment', 2)
add_para(doc, 'From an eligible rent invoice, use the M-Pesa payment action and provide the required mobile number and amount. M-Pesa STK Push is currently available only for KES invoices and uses whole-number amounts. The tenant or payer must approve the prompt on the phone by entering the M-Pesa PIN. Do not ask anyone to send their PIN to staff.')
add_heading(doc, 'Payment status and corrections', 2)
add_para(doc, 'Completed payments affect invoice balances, accounting records, and receipts. Before correcting a payment, review the invoice and receipt to understand the effect on allocations and tenant credit. Use your organisation’s authorized reversal or refund process; do not delete or recreate financial records casually.')

add_heading(doc, '5  Maintenance requests and expenses')
add_heading(doc, 'Prepare maintenance categories', 2)
add_para(doc, 'Administrators, property managers, and maintenance staff can maintain the list of maintenance categories. Use clear, consistent categories so requests and expenses can be reviewed and reported accurately.')
add_heading(doc, 'Process a staff maintenance request', 2)
for item in ['Open Maintenance Requests and create or open the request.', 'Confirm the property, unit, tenant, category, description, priority, and status fields shown in the form.', 'Record progress as work is assigned and completed.', 'Keep descriptions factual and include useful dates, observations, and resolution details.']:
    add_number(doc,item)
add_heading(doc, 'Record maintenance expenses', 2)
add_para(doc, 'Open Maintenance Expenses to record costs related to maintenance work. Link the expense to the correct related records, enter the amount and currency carefully, and retain supporting documents according to your business policy. Review the request and expense together before closing a job.')

add_heading(doc, '6  Reports and exports')
add_para(doc, 'The Reports area is available to administrators, property managers, and accountants. It consolidates invoice, payment, and maintenance information for the accessible portfolio. Use the available filters to narrow the reporting period or scope, then review the summary and detailed records before sharing results.')
add_heading(doc, 'Export a report', 2)
for item in ['Apply the filters you want to use on the Reports page.', 'Select the export action. The system exports the selected reporting scope to an Excel workbook.', 'Open the exported file and check the invoice, payment, maintenance, and summary sheets as appropriate.', 'Store or share exports only through approved channels because they may contain financial and personal data.']:
    add_number(doc,item)
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

add_heading(doc, '7  Tenant portal guide')
add_para(doc, 'The tenant portal is separate from staff operations. A tenant must have a portal account and enabled access before they can sign in. Portal users see only the information associated with their own tenancy.')
add_heading(doc, 'What tenants can do', 2)
add_table(doc, ['Portal area','Tenant action'], [
    ('Dashboard','Review the tenant-facing overview and current information.'),
    ('My Tenancy','View tenancy details supplied by the management team.'),
    ('Payments','Review payment history and open payment details where available.'),
    ('Maintenance','Submit a new request and track existing requests.'),
    ('Notifications','Read updates and mark notifications as read.'),
    ('My Profile','Update permitted personal profile details and profile photo.'),
    ('Change Password','Set a new password using the secure password form.'),
    ('Appearance','Switch between light and dark appearance from the portal interface.')], [1.65,4.9])
add_heading(doc, 'Submit a tenant maintenance request', 2)
for item in ['Open Maintenance in the portal and select the action to create a request.', 'Describe the issue clearly, including the unit area, when it started, access constraints, and any immediate safety concern.', 'Submit the request and keep the confirmation or reference available.', 'Return to the Maintenance list to see updates and open the request for details.']:
    add_number(doc,item)
add_para(doc, 'For urgent safety incidents, tenants should follow the property’s emergency procedure in addition to using the portal. The portal is a record and communication channel, not a substitute for emergency services.', 'Important:')

add_heading(doc, '8  Roles and access')
add_para(doc, 'Menus are role-based. Access to a page does not imply permission to perform every financial or administrative action. The matrix below describes the main implementation-level access groups; your organisation may also apply portfolio-level restrictions.')
add_table(doc, ['Area','Roles with access'], [
    ('User management','Administrator'),
    ('Property types','Administrator, property manager'),
    ('Owners, properties, units, tenants, leases','Administrator, property manager, reception staff, owner'),
    ('Rent invoices and payments','Administrator, accountant, owner, property manager'),
    ('Maintenance categories, requests, expenses','Administrator, property manager, maintenance staff'),
    ('Reports and Excel export','Administrator, property manager, accountant'),
    ('Settings, profile, support, notifications','All authenticated users'),
    ('Tenant portal','Tenant role only')], [2.55,4.0])

add_heading(doc, '9  Safe working practices and troubleshooting')
add_heading(doc, 'Protect records and privacy', 2)
for item in ['Work under your own account. Do not share passwords or leave a signed-in session unattended.', 'Verify the property, unit, tenant, lease, currency, amount, and billing period before saving financial records.', 'Use notes and descriptions to record facts, not passwords, PINs, or unnecessary sensitive data.', 'Download receipts and reports only to approved storage locations.', 'Use notifications and support requests for traceable follow-up.']:
    add_bullet(doc,item)
add_heading(doc, 'Common situations', 2)
add_table(doc, ['Situation','What to do'], [
    ('A menu or button is not visible','Confirm your role and sign in again. If it remains unavailable, ask an administrator to check your access.'),
    ('Invoice cannot be edited, cancelled, or deleted','Check for payments or tenant-credit allocations. Financial activity protects the record; use the approved correction process.'),
    ('M-Pesa action fails or is unavailable','Confirm that the invoice is KES, has an outstanding balance, the amount is a whole number, and the mobile number is correct.'),
    ('Duplicate invoice warning','Do not create another invoice. Locate the existing invoice for that lease and billing period.'),
    ('Tenant cannot access the portal','Open the tenant record and confirm a portal account exists and portal access is enabled.'),
    ('Need help with an error','Open Support, select the right category, include what you were doing, the record reference, and the message shown. Do not include passwords or PINs.')], [2.1,4.45])

add_heading(doc, '10  Reference checklist')
add_heading(doc, 'Monthly operations checklist', 2)
for item in ['Confirm active leases and valid monthly rents before billing.', 'Create invoices once per active lease and billing period.', 'Review due dates and outstanding balances.', 'Record or reconcile received payments, then issue or download receipts.', 'Review open maintenance requests and related costs.', 'Run and save the required management reports.']:
    add_bullet(doc,item)
add_heading(doc, 'When submitting a support request', 2)
add_para(doc, 'Include the page or module name, record type and reference number, what you expected to happen, what actually happened, the exact on-screen message, and the time of the issue. This lets support investigate without requesting sensitive credentials.')

# Footer
for section in doc.sections:
    footer = section.footer
    p = footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, 'Real Estate Management System User Manual  |  Version 1.0  |  October 2026', 8.5, False, '5A6772')

doc.core_properties.title = 'Real Estate Management System User Manual'
doc.core_properties.subject = 'End-user guide for the Real Estate Management System'
doc.core_properties.author = 'Real Estate Management System'
doc.save(OUT)
print(OUT)
